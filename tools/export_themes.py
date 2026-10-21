"""把「VCPulse題材資料庫」Excel 轉成題材雷達用的 JSON。

用法：
    python export_themes.py VCPulse題材資料庫_v11.xlsx theme-evidence.json

每次更新 Excel 後重跑一次即可。程式會檢查欄位，有問題會列出警告，但仍輸出檔案。
"""
import json
import re
import sys
from datetime import date

import openpyxl

GRADES = ["A", "B", "C", "D", "待核"]
EXPOSURES = ["高", "中", "低", "未揭露"]

RULES = {
    # 哪些證據等級算族群成員
    "membership": {
        "formal": ["A", "B"],      # 正式成員：計入族群強度
        "watch": ["C"],            # 觀察：只顯示淡色標籤
        "excluded": ["D", "待核"],  # 不顯示在雷達上
    },
    # 計算族群強度時，各曝險程度的權重
    "exposureWeight": {"高": 1.0, "中": 0.6, "低": 0.3, "未揭露": 0.5},
    # 資料日期超過幾天就標示「需更新」
    "staleDays": 180,
    # 題材成熟度：正式成員中 A 級比例
    "maturity": {
        "revenueThreshold": 0.5,   # A 級占正式成員 ≥ 50% → 有營收
        "labels": {"revenue": "有營收", "building": "布局中", "watching": "觀察中"},
    },
}

EVIDENCE_COLS = ["theme", "code", "name", "market", "role", "grade", "evidence", "exposure",
                 "exposureNote", "path", "vsOriginal", "sourceTitle", "sourceUrl", "sourceDate",
                 "verified", "manualNote"]


def norm_date(v):
    """把 2026-10-21、2026-10、2026 這類日期統一成字串。"""
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    warnings = []

    # ── 證據表 ──
    ws = wb["證據表"]
    stocks, themes = {}, {}
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        r = dict(zip(EVIDENCE_COLS, [("" if v is None else v) for v in row]))
        code = str(r["code"]).strip()
        theme = str(r["theme"]).strip()
        grade = str(r["grade"]).strip()
        exposure = str(r["exposure"]).strip()

        if not re.fullmatch(r"\d{4,6}[A-Z]?", code):
            warnings.append(f"第 {i} 列：代號「{code}」格式不對")
        if grade not in GRADES:
            warnings.append(f"第 {i} 列：證據等級「{grade}」不在 {GRADES}")
        if exposure not in EXPOSURES:
            warnings.append(f"第 {i} 列：曝險程度「{exposure}」不在 {EXPOSURES}，先當成未揭露")
            exposure = "未揭露"
        if grade in ("A", "B", "C") and not r["sourceUrl"]:
            warnings.append(f"第 {i} 列：{code} {r['name']} 是 {grade} 級但沒有來源網址")

        group, _, sub = theme.partition("：")
        themes.setdefault(theme, {"id": theme, "group": group, "name": sub or group,
                                  "status": "已建", "whitepaper": "", "policy": ""})

        s = stocks.setdefault(code, {"name": r["name"], "market": r["market"], "tags": []})
        s["tags"].append({
            "theme": theme,
            "grade": grade,
            "exposure": exposure,
            "role": r["role"],
            "evidence": r["evidence"],
            "exposureNote": r["exposureNote"],
            "path": r["path"],
            "source": {"title": r["sourceTitle"], "url": r["sourceUrl"],
                       "date": norm_date(r["sourceDate"])},
            "verified": r["verified"],
        })

    # ── 題材清單：補上白皮書出處、政策，以及還沒建證據表的題材 ──
    t = wb["題材清單"]
    last_wp = {}
    for row in t.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        group, sub, wp, _, _, policy, _, status = [("" if v is None else str(v)) for v in row[:8]]
        if wp.startswith("同上"):  # Excel 裡的「同上」換成上一列的出處
            wp = last_wp.get(group, "") + wp[2:]
        last_wp[group] = wp
        # 題材清單的子題材有時已含大題材前綴（例如「AI 伺服器：組裝代工」）
        key = sub if "：" in sub else f"{group}：{sub}"
        match = themes.get(key)
        if match is None:  # 名稱略有不同時，用包含關係找
            short = sub.split("：")[-1].split("（")[0]
            match = next((v for k, v in themes.items() if short and short in k), None)
        if match:
            match["whitepaper"] = match["whitepaper"] or wp
            match["policy"] = match["policy"] or policy
        elif status != "已建":
            themes[key] = {"id": key, "group": key.split("：")[0], "name": key.split("：")[-1],
                           "status": "未建", "whitepaper": wp, "policy": policy}

    data = {
        "version": src.rsplit("_", 1)[-1].replace(".xlsx", ""),
        "generated": date.today().isoformat(),
        "rules": RULES,
        "themes": sorted(themes.values(), key=lambda x: (x["group"], x["name"])),
        "stocks": dict(sorted(stocks.items())),
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

    n_tags = sum(len(s["tags"]) for s in stocks.values())
    built = sum(1 for x in themes.values() if x["status"] == "已建")
    print(f"輸出 {out}：{len(stocks)} 檔股票、{n_tags} 筆標籤、{built} 個已建題材、{len(themes) - built} 個未建題材")
    for w in warnings:
        print("警告：", w)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("用法：python export_themes.py 題材資料庫.xlsx 輸出.json")
    main(sys.argv[1], sys.argv[2])
